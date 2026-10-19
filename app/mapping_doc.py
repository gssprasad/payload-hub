"""Read mapping specification documents (Excel or CSV) into rows and agent-readable text."""
import csv
import io
from datetime import date, datetime

from openpyxl import load_workbook

MAX_TEXT = 150_000
MAX_ROWS_PER_SHEET = 2_000


class MappingDocError(Exception):
    """Raised when an uploaded mapping document cannot be read."""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _clean(rows: list[list[str]]) -> list[list[str]]:
    """Drop empty rows and trailing empty columns."""
    rows = [r for r in rows if any(c for c in r)]
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    return [(r + [""] * width)[:width] for r in rows]


def read(filename: str, content: bytes) -> dict:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise MappingDocError(f"Could not open the Excel file: {e}") from e
        sheets = []
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rows = _clean([[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)])
            if rows:
                sheets.append({"name": ws.title, "rows": rows})
        wb.close()
    elif name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        sheets = [{"name": "CSV", "rows": _clean([[c.strip() for c in r] for r in csv.reader(io.StringIO(text))])}]
    elif name.endswith(".xls"):
        raise MappingDocError("Old .xls files aren't supported. In Excel, use File > Save As > Excel Workbook (.xlsx).")
    else:
        raise MappingDocError("Upload an Excel workbook (.xlsx) or a CSV file.")

    sheets = [s for s in sheets if s["rows"]]
    if not sheets:
        raise MappingDocError("The document has no data.")
    for s in sheets:
        if len(s["rows"]) > MAX_ROWS_PER_SHEET:
            raise MappingDocError(f"Sheet '{s['name']}' has {len(s['rows'])} rows; the limit is {MAX_ROWS_PER_SHEET}.")

    text = to_text(sheets)
    if len(text) > MAX_TEXT:
        raise MappingDocError(
            f"The document is too large for the agent ({len(text):,} characters; limit {MAX_TEXT:,}). "
            "Remove sheets that aren't part of the mapping and upload it again."
        )
    return {"filename": filename, "sheets": sheets, "text": text}


def to_text(sheets: list[dict]) -> str:
    """Render sheets as pipe-separated tables with row numbers, so the agent can cite rows."""
    parts = []
    for s in sheets:
        lines = [f'<sheet name="{s["name"]}">']
        for i, row in enumerate(s["rows"], start=1):
            lines.append(f"{i} | " + " | ".join(c.replace("|", "/").replace("\n", " ") for c in row))
        lines.append("</sheet>")
        parts.append("\n".join(lines))
    return "\n\n".join(parts)
